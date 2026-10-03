"""
Export clean model-ready ASEAN panel data and audit reports to an Excel (.xlsx) file.
Produces:
- data/processed/asean_panel_clean.xlsx
- outputs/asean_panel_clean.xlsx
"""

from pathlib import Path
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_PROCESSED_DIR = PROJECT_ROOT / "data" / "processed"
OUTPUTS_DIR = PROJECT_ROOT / "outputs"
OUTPUTS_TABLES_DIR = OUTPUTS_DIR / "tables"

OUT_EXCEL_CLEAN = DATA_PROCESSED_DIR / "asean_panel_clean.xlsx"
OUT_EXCEL_CLEAN_OUTPUTS = OUTPUTS_DIR / "asean_panel_clean.xlsx"

def create_clean_excel():
    print(f"Loading clean datasets from {PROJECT_ROOT}...")

    clean_path = DATA_PROCESSED_DIR / "asean_panel_clean.csv"
    if not clean_path.exists():
        print(f"File not found: {clean_path}")
        return

    df_clean = pd.read_csv(clean_path)

    # Audit report
    audit_path = OUTPUTS_TABLES_DIR / "cleaning_audit_report.csv"
    df_audit = pd.read_csv(audit_path) if audit_path.exists() else pd.DataFrame()

    # Summary stats
    summary_path = OUTPUTS_TABLES_DIR / "summary_stats_overall.csv"
    df_summary = pd.read_csv(summary_path) if summary_path.exists() else pd.DataFrame()

    # Data dictionary
    dict_path = PROJECT_ROOT / "data_dictionary.csv"
    df_dict = pd.read_csv(dict_path) if dict_path.exists() else pd.DataFrame()

    # Readme Overview Table
    readme_data = [
        {"Thông tin": "Tên bộ dữ liệu", "Nội dung": "DTA301 - ASEAN Panel Data (Clean & Model-Ready)"},
        {"Thông tin": "Mục đích sử dụng", "Nội dung": "Sử dụng trực tiếp để chạy mô hình hồi quy (OLS, Fixed Effects, Random Effects, GMM)"},
        {"Thông tin": "Tỷ lệ quan sát đầy đủ", "Nội dung": "100% Complete Cases (150/150 quan sát - 10 nước x 15 năm)"},
        {"Thông tin": "Số lượng biến", "Nội dung": "60 biến (gồm 36 biến gốc sau xử lý + biến trễ L1/L2 + logarit + sai phân + biến giả)"},
        {"Thông tin": "Phương pháp xử lý khuyết", "Nội dung": "Nội suy chuỗi thời gian nội bộ từng nước (Linear Interpolation) + Trung vị nhóm nước"},
        {"Thông tin": "Phương pháp xử lý ngoại lai", "Nội dung": "Winsorization 1% - 99% đối với biến phân phối lệch (secure_servers, trade_openness)"},
        {"Thông tin": "Sheet: Panel_Clean_Model", "Nội dung": "Toàn bộ bảng dữ liệu sạch 150 dòng x 60 cột, không còn giá trị rỗng"},
        {"Thông tin": "Sheet: Cleaning_Audit", "Nội dung": "Báo cáo đối chiếu Before vs After (số lượng giá trị NaN trước và sau khi xử lý)"},
        {"Thông tin": "Sheet: Summary_Statistics", "Nội dung": "Bảng thống kê mô tả tổng thể (Mean, Std, Min, Max, P25, Median, P75, Skewness)"},
        {"Thông tin": "Sheet: Data_Dictionary", "Nội dung": "Từ điển 60 biến số kinh tế lượng với giải thích chi tiết tiếng Việt và nhãn Stata"},
    ]
    df_readme = pd.DataFrame(readme_data)

    sheets = {
        "Readme_Overview": df_readme,
        "Panel_Clean_Model": df_clean,
        "Cleaning_Audit": df_audit,
        "Summary_Statistics": df_summary,
        "Data_Dictionary": df_dict,
    }

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    regular_font = Font(name="Segoe UI", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    header_border = Border(
        left=Side(style="thin", color="1B365D"),
        right=Side(style="thin", color="1B365D"),
        top=Side(style="medium", color="1B365D"),
        bottom=Side(style="medium", color="1B365D"),
    )

    for sheet_name, df in sheets.items():
        if df.empty:
            continue
        ws = wb.create_sheet(title=sheet_name)
        ws.views.sheetView[0].showGridLines = True

        headers = list(df.columns)
        ws.append(headers)
        ws.row_dimensions[1].height = 28

        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = header_border

        records = df.where(pd.notnull(df), None).values.tolist()
        for row_idx, row_val in enumerate(records, start=2):
            ws.append(row_val)
            ws.row_dimensions[row_idx].height = 20
            for col_idx, val in enumerate(row_val, start=1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = regular_font
                cell.border = thin_border
                if isinstance(val, (int, float)):
                    if isinstance(val, int) or (isinstance(val, float) and val.is_integer()):
                        cell.number_format = '#,##0'
                    else:
                        cell.number_format = '#,##0.0000'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif isinstance(val, str) and len(val) <= 10 and (val.isupper() or val.isdigit()):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        if sheet_name == "Panel_Clean_Model":
            ws.freeze_panes = "D2"
        else:
            ws.freeze_panes = "A2"

        for col_idx, col in enumerate(ws.columns, start=1):
            max_len = 0
            sample_cells = [col[0]] + list(col[1:min(100, len(col))])
            for cell in sample_cells:
                if cell.value is not None:
                    s_len = len(str(cell.value))
                    if s_len > max_len:
                        max_len = s_len
            col_letter = get_column_letter(col_idx)
            calc_width = max(max_len + 4, 12)
            ws.column_dimensions[col_letter].width = min(calc_width, 60)

    OUT_EXCEL_CLEAN.parent.mkdir(parents=True, exist_ok=True)
    OUT_EXCEL_CLEAN_OUTPUTS.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_EXCEL_CLEAN)
    wb.save(OUT_EXCEL_CLEAN_OUTPUTS)
    print(f"Saved clean excel workbook to: {OUT_EXCEL_CLEAN} and {OUT_EXCEL_CLEAN_OUTPUTS}")

if __name__ == "__main__":
    create_clean_excel()
